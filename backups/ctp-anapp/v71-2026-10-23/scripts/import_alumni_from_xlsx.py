#!/usr/bin/env python3
"""Reconcile the reviewed CTP ANAPP alumni workbook with Supabase.

Requires SUPABASE_ACCESS_TOKEN in the process environment. No credential is
written to disk or printed.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import secrets
import subprocess
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

PROJECT_REF = "egihufoowzfykieozusr"
PROJECT_URL = f"https://{PROJECT_REF}.supabase.co"
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def normalize_email(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def is_yes(value: Any) -> bool:
    return str(value or "").strip().upper() in {"SIM", "S", "YES", "TRUE", "1"}


def get_service_key(access_token: str) -> str:
    env = os.environ.copy()
    env["SUPABASE_ACCESS_TOKEN"] = access_token
    result = subprocess.run(
        [
            "npx",
            "supabase",
            "projects",
            "api-keys",
            "--project-ref",
            PROJECT_REF,
            "--output",
            "json",
        ],
        check=True,
        capture_output=True,
        text=True,
        env=env,
        timeout=60,
    )
    keys = json.loads(result.stdout)
    return next(item["api_key"] for item in keys if item["id"] == "service_role")


def request(
    service_key: str,
    path: str,
    *,
    method: str = "GET",
    body: Any | None = None,
    prefer: str | None = None,
) -> tuple[Any, dict[str, str], int]:
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    data = None if body is None else json.dumps(body).encode()
    req = urllib.request.Request(
        f"{PROJECT_URL}{path}",
        data=data,
        headers=headers,
        method=method,
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            raw = response.read()
            parsed = json.loads(raw) if raw else None
            return parsed, dict(response.headers), response.status
    except urllib.error.HTTPError as error:
        details = error.read().decode(errors="ignore")
        raise RuntimeError(f"HTTP {error.code} em {path}: {details[:500]}") from error


def load_records(path: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    column = {name: index for index, name in enumerate(headers)}

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    invalid: list[int] = []

    for row_number, values in enumerate(rows, start=2):
        email = normalize_email(values[column["E-mail"]])
        if not EMAIL_RE.match(email):
            invalid.append(row_number)
            continue
        record = {
            "source_row": row_number,
            "email": email,
            "full_name": clean_text(values[column["NOME"]]),
            "cpf": clean_text(values[column["CPF"]]),
            "phone": clean_text(values[column["Telefone"]]) if "Telefone" in column else None,
            "certified": is_yes(values[column["Certificado Emitido"]]),
            "enrolled_at": values[column["Data de Inscrição"]],
        }
        grouped[email].append(record)

    result: list[dict[str, Any]] = []
    duplicate_groups: list[list[int]] = []
    for email, items in grouped.items():
        if len(items) > 1:
            duplicate_groups.append([item["source_row"] for item in items])

        def rank(item: dict[str, Any]) -> tuple[int, str, int]:
            completeness = sum(bool(item.get(key)) for key in ("full_name", "cpf", "phone"))
            enrolled = item.get("enrolled_at")
            if isinstance(enrolled, (datetime, date)):
                date_key = enrolled.isoformat()
            else:
                date_key = str(enrolled or "")
            return completeness, date_key, item["source_row"]

        selected = max(items, key=rank).copy()
        selected["alumni_type"] = (
            "certified_professional"
            if any(item["certified"] for item in items)
            else "alumni"
        )
        selected.pop("certified", None)
        selected.pop("enrolled_at", None)
        result.append(selected)

    result.sort(key=lambda item: item["email"])
    return result, {
        "source_rows": sheet.max_row - 1,
        "valid_unique_emails": len(result),
        "invalid_rows": invalid,
        "duplicate_groups": duplicate_groups,
    }


def list_auth_users(service_key: str) -> list[dict[str, Any]]:
    users: list[dict[str, Any]] = []
    page = 1
    while True:
        data, _, _ = request(
            service_key,
            f"/auth/v1/admin/users?page={page}&per_page=1000",
        )
        batch = data.get("users", []) if isinstance(data, dict) else []
        users.extend(batch)
        if len(batch) < 1000:
            return users
        page += 1


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    parser.add_argument("--report", default="artifacts/private/alumni-import-result.json")
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--workers", type=int, default=16)
    args = parser.parse_args()

    access_token = os.environ.get("SUPABASE_ACCESS_TOKEN")
    if not access_token:
        raise SystemExit("SUPABASE_ACCESS_TOKEN ausente.")

    records, workbook_report = load_records(args.workbook)
    service_key = get_service_key(access_token)

    roles, _, _ = request(
        service_key,
        "/rest/v1/user_roles?select=user_id,role,user_origin&limit=5000",
    )
    protected_admin_ids = {
        item["user_id"]
        for item in roles
        if item.get("role") in {"admin", "master_admin"}
    }
    legacy_ids = {
        item["user_id"]
        for item in roles
        if item.get("user_origin") in {"alumni", "certified_professional", "alumni_certified"}
    }
    whitelist, _, _ = request(
        service_key,
        "/rest/v1/alumni_whitelist?select=id,email,linked_user_id&limit=5000",
    )
    legacy_ids.update(
        item["linked_user_id"]
        for item in whitelist
        if item.get("linked_user_id")
    )

    deleted_auth = 0
    if not args.resume:
        for user_id in sorted(legacy_ids - protected_admin_ids):
            request(service_key, f"/auth/v1/admin/users/{user_id}", method="DELETE")
            deleted_auth += 1

        request(
            service_key,
            "/rest/v1/alumni_whitelist?id=not.is.null",
            method="DELETE",
            prefer="return=minimal",
        )

    auth_users = list_auth_users(service_key)
    users_by_email = {
        normalize_email(item.get("email")): item for item in auth_users if item.get("email")
    }

    created = 0
    reused = 0
    granted = 0
    certified = 0
    errors: list[dict[str, Any]] = []

    def import_record(record: dict[str, Any]) -> dict[str, Any]:
        try:
            request(
                service_key,
                "/rest/v1/rpc/admin_import_alumni",
                method="POST",
                body={
                    "p_email": record["email"],
                    "p_full_name": record["full_name"],
                    "p_cpf": record["cpf"],
                    "p_phone": record["phone"],
                    "p_alumni_type": record["alumni_type"],
                    "p_cert_number": None,
                    "p_year": None,
                },
            )

            user = users_by_email.get(record["email"])
            if not user:
                user, _, _ = request(
                    service_key,
                    "/auth/v1/admin/users",
                    method="POST",
                    body={
                        "email": record["email"],
                        "password": secrets.token_urlsafe(32) + "Aa1!",
                        "email_confirm": True,
                        "user_metadata": {
                            "full_name": record["full_name"],
                            "cpf": record["cpf"],
                            "user_origin": record["alumni_type"],
                        },
                    },
                )
                was_created = True
            else:
                was_created = False

            grant, _, _ = request(
                service_key,
                "/rest/v1/rpc/check_alumni_and_grant_access",
                method="POST",
                body={"p_user_id": user["id"], "p_email": record["email"]},
            )
            return {
                "created": was_created,
                "granted": bool(grant and grant.get("access_granted_without_payment")),
                "certified": record["alumni_type"] == "certified_professional",
            }
        except Exception as error:
            return {
                "error": {
                    "source_row": record["source_row"],
                    "email_domain": record["email"].split("@")[-1],
                    "error": str(error)[:500],
                },
            }

    with ThreadPoolExecutor(max_workers=max(1, min(args.workers, 24))) as executor:
        futures = [executor.submit(import_record, record) for record in records]
        for future in as_completed(futures):
            result = future.result()
            if result.get("error"):
                errors.append(result["error"])
                continue
            if result["created"]:
                created += 1
            else:
                reused += 1
            granted += int(result["granted"])
            certified += int(result["certified"])

    def exact_count(path: str) -> int:
        _, headers, _ = request(service_key, path, prefer="count=exact")
        content_range = headers.get("Content-Range", "*/0")
        return int(content_range.rsplit("/", 1)[-1])

    verification = {
        "alumni_whitelist": exact_count(
            "/rest/v1/alumni_whitelist?select=id&limit=1"
        ),
        "profiles_is_alumni": exact_count(
            "/rest/v1/profiles?select=id&is_alumni=eq.true&limit=1"
        ),
        "linked_whitelist": exact_count(
            "/rest/v1/alumni_whitelist?select=id&linked_user_id=not.is.null&limit=1"
        ),
        "student_roles_linked": exact_count(
            "/rest/v1/user_roles?select=id&alumni_whitelist_id=not.is.null&limit=1"
        ),
    }

    report = {
        "workbook": workbook_report,
        "deleted_previous_auth_users": deleted_auth,
        "created_auth_users": created,
        "reused_auth_users": reused,
        "access_granted": granted,
        "certified_professionals": certified,
        "errors": errors,
        "verification": verification,
    }
    os.makedirs(os.path.dirname(args.report), exist_ok=True)
    with open(args.report, "w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
    os.chmod(args.report, 0o600)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors and verification["linked_whitelist"] == len(records) else 1


if __name__ == "__main__":
    sys.exit(main())
